"""
Generate an A3 Portrait Academic Poster for MCA Project Expo:
"Cyber Threat Intelligence Pipeline with NLP & LLM"
Matching the exact blueprint structure requested:
1. Title & surrounding icons
2. Problem Statement (highlight card / note)
3. Architecture (Dashboard UI mockup + Modular System Flowchart)
4. Steps Taken to Execute CTI Pipeline (2 columns of technical steps)
5. Bottom Category Icons & Student Attribution
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_a3_poster(output_path="CTI_Poster_A3.pptx"):
    prs = Presentation()
    
    # Set slide dimensions to A3 Portrait (11.693 x 16.535 inches)
    prs.slide_width = Inches(11.693)
    prs.slide_height = Inches(16.535)
    
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    
    # Colors
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    DARK_NAVY = RGBColor(0x0A, 0x19, 0x2F)
    CYBER_BLUE = RGBColor(0x00, 0x78, 0xD4)
    ACCENT_CYAN = RGBColor(0x00, 0xA8, 0xE8)
    TEXT_DARK = RGBColor(0x1F, 0x29, 0x37)
    TEXT_MUTED = RGBColor(0x4B, 0x55, 0x63)
    CARD_BG = RGBColor(0xF8, 0xFA, 0xFC)
    NOTE_BG = RGBColor(0xFE, 0xF9, 0xC3) # Post-it / taped note yellow
    BORDER_COLOR = RGBColor(0xCB, 0xD5, 0xE1)
    ACCENT_GREEN = RGBColor(0x10, 0xB9, 0x81)
    ACCENT_RED = RGBColor(0xEF, 0x44, 0x44)
    
    # Outer Border Frame
    border = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(0.4), Inches(10.893), Inches(15.735)
    )
    border.fill.solid()
    border.fill.fore_color.rgb = WHITE
    border.line.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    border.line.width = Pt(2.5)
    
    # Header Attribution Banner
    header_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.55), Inches(10.493), Inches(0.4))
    tf_h = header_box.text_frame
    tf_h.word_wrap = True
    p_h = tf_h.paragraphs[0]
    p_h.alignment = PP_ALIGN.CENTER
    r_h = p_h.add_run()
    r_h.text = "ADMINISTRATIVE MANAGEMENT COLLEGE  |  DEPARTMENT OF MCA  |  PROJECT EXPO 2026"
    r_h.font.size = Pt(11)
    r_h.font.bold = True
    r_h.font.color.rgb = CYBER_BLUE

    # Top Corner Icons / Badges
    icons_top = [
        ("🛡️ SHIELD", Inches(0.8), Inches(1.15)),
        ("📡 RADAR", Inches(2.7), Inches(1.1)),
        ("💻 TERMINAL", Inches(8.0), Inches(1.1)),
        ("🔒 SECURE", Inches(9.8), Inches(1.15))
    ]
    for text, left, top in icons_top:
        tb = slide.shapes.add_textbox(left, top, Inches(1.6), Inches(0.35))
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = TEXT_MUTED

    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.45), Inches(10.493), Inches(1.1))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.alignment = PP_ALIGN.CENTER
    r_t = p_t.add_run()
    r_t.text = "CYBER THREAT INTELLIGENCE"
    r_t.font.size = Pt(32)
    r_t.font.bold = True
    r_t.font.color.rgb = DARK_NAVY
    
    p_sub = tf_t.add_paragraph()
    p_sub.alignment = PP_ALIGN.CENTER
    r_sub = p_sub.add_run()
    r_sub.text = "Automated NLP & LLM Pipeline for Threat Extraction, MITRE ATT&CK Mapping & Knowledge Graphs"
    r_sub.font.size = Pt(14)
    r_sub.font.bold = True
    r_sub.font.color.rgb = CYBER_BLUE

    # Student Info Banner
    info_box = slide.shapes.add_textbox(Inches(0.6), Inches(2.55), Inches(10.493), Inches(0.35))
    tf_i = info_box.text_frame
    p_i = tf_i.paragraphs[0]
    p_i.alignment = PP_ALIGN.CENTER
    r_i = p_i.add_run()
    r_i.text = "Presenter: Siddaray Bhajabalakar  |  USN: P03AC24S126056  |  MCA Final Year"
    r_i.font.size = Pt(12)
    r_i.font.bold = True
    r_i.font.color.rgb = TEXT_DARK

    # ─────────────────────────────────────────────────────────────
    # SECTION 1: PROBLEM STATEMENT (Card styled like blueprint note)
    # ─────────────────────────────────────────────────────────────
    ps_card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.2), Inches(3.0), Inches(7.293), Inches(2.1)
    )
    ps_card.fill.solid()
    ps_card.fill.fore_color.rgb = NOTE_BG
    ps_card.line.color.rgb = RGBColor(0xCA, 0x8A, 0x04)
    ps_card.line.width = Pt(1.5)

    # Pin / Tape Graphic on Problem Statement
    tape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(5.1), Inches(2.9), Inches(1.5), Inches(0.22)
    )
    tape.fill.solid()
    tape.fill.fore_color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
    tape.line.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    tape.line.width = Pt(0.75)

    ps_tb = slide.shapes.add_textbox(Inches(2.4), Inches(3.15), Inches(6.9), Inches(1.8))
    tf_ps = ps_tb.text_frame
    tf_ps.word_wrap = True
    p_pst = tf_ps.paragraphs[0]
    p_pst.alignment = PP_ALIGN.CENTER
    r_pst = p_pst.add_run()
    r_pst.text = "📌 PROBLEM STATEMENT"
    r_pst.font.size = Pt(16)
    r_pst.font.bold = True
    r_pst.font.color.rgb = RGBColor(0x9A, 0x34, 0x12)

    problems = [
        "Unstructured Intelligence Overload: Analysts manually read hundreds of pages of raw threat reports.",
        "Delayed Threat Detection: Manual extraction of IPs, Hashes, and CVEs causes critical lag in SOC response.",
        "Disconnected Threat Context: Isolated indicators lack correlation to Threat Actors, TTPs, and Malware.",
        "High Alert Fatigue: Security analysts spend 60%+ time on routine entity parsing instead of remediation."
    ]
    for prob in problems:
        p = tf_ps.add_paragraph()
        p.space_before = Pt(3)
        r = p.add_run()
        r.text = "•  " + prob
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x1C, 0x19, 0x17)

    # ─────────────────────────────────────────────────────────────
    # SECTION 2: ARCHITECTURE (Left: Dashboard, Right: Flow Diagram)
    # ─────────────────────────────────────────────────────────────
    arch_title = slide.shapes.add_textbox(Inches(0.6), Inches(5.25), Inches(10.493), Inches(0.45))
    tf_at = arch_title.text_frame
    p_at = tf_at.paragraphs[0]
    p_at.alignment = PP_ALIGN.CENTER
    r_at = p_at.add_run()
    r_at.text = "⚙️ ARCHITECTURE & SYSTEM DESIGN"
    r_at.font.size = Pt(20)
    r_at.font.bold = True
    r_at.font.color.rgb = DARK_NAVY

    # Left Box: Security Analyst Interface / Dashboard Mockup
    dash_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(4.5), Inches(4.4)
    )
    dash_box.fill.solid()
    dash_box.fill.fore_color.rgb = RGBColor(0x0F, 0x17, 0x2A) # Dark terminal aesthetic
    dash_box.line.color.rgb = CYBER_BLUE
    dash_box.line.width = Pt(1.5)

    dash_tb = slide.shapes.add_textbox(Inches(0.9), Inches(5.9), Inches(4.3), Inches(4.2))
    tf_d = dash_tb.text_frame
    tf_d.word_wrap = True
    
    p_dt = tf_d.paragraphs[0]
    r_dt = p_dt.add_run()
    r_dt.text = "🖥️ ANALYST CONSOLE & REST API\n"
    r_dt.font.size = Pt(13)
    r_dt.font.bold = True
    r_dt.font.color.rgb = ACCENT_CYAN

    dash_items = [
        ("Input Source", "Unstructured CTI (APT Reports, MISP Feeds)"),
        ("Active Feed", "APT29 Cozy Bear / Lazarus Campaign Report"),
        ("Extracted IOCs", "IP: 192.168.1.50 | SHA256: 4a3f8b... | CVE-2024-21413"),
        ("MITRE Tagged", "T1059 (Command & Script) | T1566 (Phishing)"),
        ("KG Relations", "APT29 ──[uses]──> Backdoor.CobaltStrike"),
        ("LLM Threat Brief", "\"High risk: Nation-state actor targeting energy sector\""),
        ("API Endpoints", "POST /analyze | GET /threat_actor | GET /knowledge_graph")
    ]
    for label, val in dash_items:
        p = tf_d.add_paragraph()
        p.space_before = Pt(4)
        r1 = p.add_run()
        r1.text = f"[{label}]: "
        r1.font.size = Pt(9.5)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8)
        r2 = p.add_run()
        r2.text = val
        r2.font.size = Pt(9)
        r2.font.color.rgb = WHITE

    # Right Box: Hierarchical Modular Flowchart (Matching Blueprint Tree)
    # Root Node
    root_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.0), Inches(5.8), Inches(4.6), Inches(0.65)
    )
    root_shape.fill.solid()
    root_shape.fill.fore_color.rgb = CYBER_BLUE
    root_shape.line.color.rgb = DARK_NAVY
    tf_r = root_shape.text_frame
    tf_r.word_wrap = True
    p_r = tf_r.paragraphs[0]
    p_r.alignment = PP_ALIGN.CENTER
    r_r = p_r.add_run()
    r_r.text = "CTI ANALYSIS ENGINE (PIPELINE CORE)"
    r_r.font.size = Pt(13)
    r_r.font.bold = True
    r_r.font.color.rgb = WHITE

    # Level 1 Modules (3 Blocks across)
    l1_modules = [
        ("NLP & NER\nExtractor", Inches(5.6), Inches(6.9), Inches(1.5)),
        ("MITRE ATT&CK\nClassifier", Inches(7.35), Inches(6.9), Inches(1.6)),
        ("Knowledge Graph\n& LLM Engine", Inches(9.2), Inches(6.9), Inches(1.6))
    ]
    for title, left, top, width in l1_modules:
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.75))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor(0xE0, 0xF2, 0xFE)
        sh.line.color.rgb = CYBER_BLUE
        sh.line.width = Pt(1.2)
        tf = sh.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = title
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = DARK_NAVY

    # Level 2 Sub-Modules (3 Blocks underneath)
    l2_modules = [
        ("Regex Patterns\n+ spaCy Models\n(IP, Hash, CVE)", Inches(5.6), Inches(8.15), Inches(1.5)),
        ("100+ TTPs Database\nTactics TA0001-TA0043\nConfidence Scoring", Inches(7.35), Inches(8.15), Inches(1.6)),
        ("NetworkX Graph\nEntity Linking\nLLM Summarizer", Inches(9.2), Inches(8.15), Inches(1.6))
    ]
    for desc, left, top, width in l2_modules:
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.95))
        sh.fill.solid()
        sh.fill.fore_color.rgb = WHITE
        sh.line.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        sh.line.width = Pt(1)
        tf = sh.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = desc
        r.font.size = Pt(8.5)
        r.font.color.rgb = TEXT_DARK

    # Level 3 Bottom Output Module
    out_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.6), Inches(9.45), Inches(5.2), Inches(0.75)
    )
    out_shape.fill.solid()
    out_shape.fill.fore_color.rgb = RGBColor(0xEC, 0xFD, 0xF5)
    out_shape.line.color.rgb = ACCENT_GREEN
    out_shape.line.width = Pt(1.5)
    tf_o = out_shape.text_frame
    p_o = tf_o.paragraphs[0]
    p_o.alignment = PP_ALIGN.CENTER
    r_o = p_o.add_run()
    r_o.text = "INTEGRATION LAYER: Flask REST API + Docker Microservice"
    r_o.font.size = Pt(11)
    r_o.font.bold = True
    r_o.font.color.rgb = RGBColor(0x06, 0x5F, 0x46)

    # ─────────────────────────────────────────────────────────────
    # SECTION 3: STEPS TAKEN TO EXECUTE (Bottom Container - 2 Columns)
    # ─────────────────────────────────────────────────────────────
    steps_container = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(10.5), Inches(10.093), Inches(4.3)
    )
    steps_container.fill.solid()
    steps_container.fill.fore_color.rgb = CARD_BG
    steps_container.line.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    steps_container.line.width = Pt(1.8)

    # Steps Container Title
    sc_title = slide.shapes.add_textbox(Inches(1.0), Inches(10.6), Inches(9.7), Inches(0.4))
    tf_sct = sc_title.text_frame
    p_sct = tf_sct.paragraphs[0]
    p_sct.alignment = PP_ALIGN.CENTER
    r_sct = p_sct.add_run()
    r_sct.text = "🚀 STEPS TAKEN TO EXECUTE CTI PIPELINE"
    r_sct.font.size = Pt(18)
    r_sct.font.bold = True
    r_sct.font.color.rgb = DARK_NAVY

    # Left Column (Steps 1 to 3)
    col1_tb = slide.shapes.add_textbox(Inches(1.0), Inches(11.1), Inches(4.7), Inches(3.2))
    tf_c1 = col1_tb.text_frame
    tf_c1.word_wrap = True
    
    col1_steps = [
        ("Phase 1: Requirements & Threat Data Collection",
         "Surveyed SOC bottlenecks; ingested CTI corpora from MISP feeds, CTI-HAL benchmark dataset, and public cyber advisories to standardize raw input schemas."),
        ("Phase 2: Modular Pipeline & NER Architecture",
         "Architected decoupled microservices for ingestion, feature extraction, entity linking, and queryable persistence with robust exception handling."),
        ("Phase 3: Automated IOC Extraction Engine",
         "Engineered hybrid NLP regex and spaCy tokenizers extracting IPv4/IPv6, SHA256/MD5 hashes, URLs, CVE numbers, malware names, and threat actors.")
    ]
    for i, (stitle, sdesc) in enumerate(col1_steps):
        p1 = tf_c1.add_paragraph() if i > 0 else tf_c1.paragraphs[0]
        p1.space_before = Pt(6)
        r1 = p1.add_run()
        r1.text = f"• {stitle}\n"
        r1.font.size = Pt(10.5)
        r1.font.bold = True
        r1.font.color.rgb = CYBER_BLUE
        
        r2 = p1.add_run()
        r2.text = f"  {sdesc}"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_DARK

    # Right Column (Steps 4 to 6)
    col2_tb = slide.shapes.add_textbox(Inches(5.9), Inches(11.1), Inches(4.7), Inches(3.2))
    tf_c2 = col2_tb.text_frame
    tf_c2.word_wrap = True
    
    col2_steps = [
        ("Phase 4: MITRE ATT&CK TTP Classification",
         "Integrated 100+ Enterprise ATT&CK techniques, tagging reports with Tactics (TA0001–TA0043) and Techniques (T1059, T1566) with confidence scores."),
        ("Phase 5: Knowledge Graph & LLM Reasoning",
         "Constructed dynamic NetworkX graph modeling Actor-Uses-Malware-Targets-Sector graphs; integrated LLM for automated threat brief generation and Q&A."),
        ("Phase 6: REST API, Dockerization & Validation",
         "Developed high-throughput Flask API endpoints (/analyze, /query), containerized with Docker, and evaluated on APT29/Lazarus benchmark reports.")
    ]
    for i, (stitle, sdesc) in enumerate(col2_steps):
        p1 = tf_c2.add_paragraph() if i > 0 else tf_c2.paragraphs[0]
        p1.space_before = Pt(6)
        r1 = p1.add_run()
        r1.text = f"• {stitle}\n"
        r1.font.size = Pt(10.5)
        r1.font.bold = True
        r1.font.color.rgb = CYBER_BLUE
        
        r2 = p1.add_run()
        r2.text = f"  {sdesc}"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_DARK

    # Bottom Icons Row (Matching Care Connect blueprint: First Aid, Puzzle, Clipboard, Mobile)
    # In CTI: Terminal, Knowledge Graph, Shield, Cloud Server
    bottom_badges = [
        ("💻 Terminal CLI / API", Inches(1.2), Inches(14.9), Inches(2.0)),
        ("🕸️ Knowledge Graph", Inches(3.5), Inches(14.9), Inches(2.0)),
        ("🛡️ MITRE ATT&CK", Inches(5.8), Inches(14.9), Inches(2.0)),
        ("🐳 Docker Container", Inches(8.1), Inches(14.9), Inches(2.0))
    ]
    for btext, bleft, btop, bwidth in bottom_badges:
        bshape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bleft, btop, bwidth, Inches(0.42))
        bshape.fill.solid()
        bshape.fill.fore_color.rgb = WHITE
        bshape.line.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        bshape.line.width = Pt(1)
        tf = bshape.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = btext
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = DARK_NAVY

    # Footer note
    footer_box = slide.shapes.add_textbox(Inches(0.6), Inches(15.75), Inches(10.493), Inches(0.3))
    tf_ft = footer_box.text_frame
    p_ft = tf_ft.paragraphs[0]
    p_ft.alignment = PP_ALIGN.CENTER
    r_ft = p_ft.add_run()
    r_ft.text = "Department of MCA  •  Administrative Management College  •  A3 Poster Presentation for Project Expo 2026"
    r_ft.font.size = Pt(9)
    r_ft.font.color.rgb = TEXT_MUTED

    prs.save(output_path)
    print(f"A3 Poster saved successfully to: {output_path}")

if __name__ == "__main__":
    create_a3_poster()
